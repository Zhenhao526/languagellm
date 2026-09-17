"""Build this project's Chinese research report with the bundled Python runtime.

No document is created on import. Run only after the required artifact-start mark.
The source is deliberately a small Markdown subset: headings, paragraphs, lists,
block quotes, fenced code, links, emphasis, tables, local images, and page breaks.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import unicodedata
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
BUNDLED_PYTHON = Path('/Users/xia/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3.12')
DEFAULT_OUTPUT = HERE / '预训练视觉主体中符号约定的形成与适应.docx'
BODY_FONT = 'Songti SC'
HEADING_FONT = 'PingFang SC'
LATIN_FONT = 'Cambria'
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
LEFT_MARGIN_CM = RIGHT_MARGIN_CM = 2.2
CONTENT_WIDTH_CM = PAGE_WIDTH_CM - LEFT_MARGIN_CM - RIGHT_MARGIN_CM
BLACK = '000000'
CAPTION_RE = re.compile(r'^(?:图|表)\s*\d+\s*[：:．.、\s]|^(?:Figure|Table)\s+\d+[.:\s]', re.I)


def set_font(target, *, size=11, bold=False, east_asia=BODY_FONT, latin=LATIN_FONT, color=BLACK):
    target.font.name = latin
    target.font.size = Pt(size)
    target.font.bold = bold
    target.font.color.rgb = RGBColor.from_string(color)
    rpr = target._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for field in ('ascii', 'hAnsi', 'cs'):
        fonts.set(qn(f'w:{field}'), latin)
    fonts.set(qn('w:eastAsia'), east_asia)
    for field in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn(f'w:{field}'), None)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'zh-CN')


def p_settings(p, *, size=11, after=5, before=0, line=1.32, keep_next=False):
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    # Songti's ascender metrics vary across Word/LibreOffice. An exact, generous
    # baseline distance avoids a renderer-dependent doubling of Chinese leading.
    f.line_spacing = Pt(max(size * line, size + 4.5))
    f.widow_control = True
    f.keep_with_next = keep_next
    f.keep_together = False
    ppr = p._p.get_or_add_pPr()
    for name, value in [('w:kinsoku', '1'), ('w:wordWrap', '1'), ('w:overflowPunct', '0'), ('w:snapToGrid', '0')]:
        element = ppr.find(qn(name))
        if element is None:
            element = OxmlElement(name)
            ppr.append(element)
        element.set(qn('w:val'), value)


def breakable_text(value):
    """Allow long file names / URLs to wrap without changing their visible text."""
    return re.sub(r'([/_=?:&])', lambda match: match[1] + '\u200b', value) if len(value) > 45 else value


def hyperlink(p, label, target, *, size=11, bold=False):
    h = OxmlElement('w:hyperlink')
    if target.startswith('#'):
        h.set(qn('w:anchor'), re.sub(r'\W', '_', target[1]))
    else:
        if target.startswith('/'):
            # Relative links survive moving the report and its surrounding project.
            from os.path import relpath
            target = relpath(unquote(target), HERE)
        h.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    h.set(qn('w:history'), '1')
    run = p.add_run(breakable_text(label))
    set_font(run, size=size, bold=bold, color='1F3D5B')
    run.font.underline = True
    h.append(run._r)
    p._p.append(h)


def parse_link(text, start):
    """Find Markdown labels/targets, including balanced parentheses in URLs."""
    close = text.find(']', start + 1)
    if close < 0 or text[close + 1:close + 2] != '(':
        return None
    pos, depth = close + 2, 1
    while pos < len(text):
        if text[pos] == '(':
            depth += 1
        elif text[pos] == ')':
            depth -= 1
            if depth == 0:
                target = text[close + 2:pos].strip()
                if target.startswith('<') and target.endswith('>'):
                    target = target[1:-1]
                return text[start + 1:close], target, pos + 1
        pos += 1
    return None


def inline(p, text, *, size=11, bold=False, italic=False, code=False, heading=False):
    pos = 0
    while pos < len(text):
        token = None
        if text[pos] == '[':
            token = parse_link(text, pos)
        if token:
            label, target, pos = token
            hyperlink(p, re.sub(r'[*`]', '', label), target, size=size, bold=bold)
            continue
        if text.startswith('**', pos):
            end = text.find('**', pos + 2)
            if end >= 0:
                inline(p, text[pos + 2:end], size=size, bold=True, italic=italic, heading=heading)
                pos = end + 2
                continue
        if text[pos] in ('*', '`'):
            delimiter = text[pos]
            end = text.find(delimiter, pos + 1)
            if end >= 0:
                inline(p, text[pos + 1:end], size=size, bold=bold,
                       italic=delimiter == '*', code=delimiter == '`', heading=heading)
                pos = end + 1
                continue
        if text.startswith('<br>', pos) or text.startswith('<br/>', pos):
            length = 4 if text.startswith('<br>', pos) else 5
            p.add_run().add_break()
            pos += length
            continue
        end = pos + 1
        while end < len(text) and text[end] not in '[*`<' and not text.startswith('**', end):
            end += 1
        chunk = text[pos:end].replace('\\|', '|')
        run = p.add_run(breakable_text(chunk))
        set_font(run, size=size, bold=bold, east_asia=HEADING_FONT if heading else BODY_FONT,
                 latin='Menlo' if code else LATIN_FONT)
        run.font.italic = italic
        pos = end


def make_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(PAGE_WIDTH_CM), Cm(PAGE_HEIGHT_CM)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2.0)
    section.left_margin, section.right_margin = Cm(LEFT_MARGIN_CM), Cm(RIGHT_MARGIN_CM)
    section.header_distance, section.footer_distance = Cm(.9), Cm(.9)
    normal = doc.styles['Normal']
    set_font(normal, size=11)
    normal.paragraph_format.line_spacing = Pt(15.5)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [('Title', 22), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
        style = doc.styles[name]
        set_font(style, size=size, bold=name != 'Subtitle', east_asia=HEADING_FONT)
        style.paragraph_format.space_before = Pt(15 if name != 'Title' else 2)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.line_spacing = Pt(size + 5)
        ppr = style._element.find(qn('w:pPr'))
        if ppr is not None:
            border = ppr.find(qn('w:pBdr'))
            if border is not None:
                ppr.remove(border)
    set_font(doc.styles['Caption'], size=9.5)
    doc.styles['Caption'].paragraph_format.space_after = Pt(7)
    doc.styles['Caption'].paragraph_format.line_spacing = 1.15
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    run = footer.add_run('1')
    set_font(run, size=9)
    field.append(run._r)
    footer._p.append(field)
    doc.core_properties.title = '预训练视觉主体中符号约定的形成与适应'
    doc.core_properties.subject = '视觉资源合作任务中的通信形成 伙伴接触与适应机制'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.keywords = '涌现通信 视觉主体 资源合作 伙伴接触 先导研究'
    return doc


def table_cells(line):
    value = line.strip().strip('|')
    return [part.strip().replace('\\|', '|') for part in re.split(r'(?<!\\)\|', value)]


def text_width(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text)
    text = re.sub(r'[*`]', '', text)
    return sum(1 if unicodedata.east_asian_width(c) in 'WF' else .52 for c in text)


def numeric_cell(text):
    return bool(re.fullmatch(r'[+−–\-]?[\d.,%％/±≈<>≤≥()（）\s]+(?:个百分点|次|对|参数|秒|分钟)?', text))


def column_widths(rows, explicit=None):
    count = len(rows[0])
    # Known report tables receive deliberate widths. This changes layout only;
    # all cell contents continue to come verbatim from the Markdown source.
    report_widths = {
        ('模块', 3): [3.0, 2.5, 11.1],
        ('条件', 5): ([2.5, 2.8, 2.8, 3.0, 5.5] if rows[0][1].startswith('前')
                      else [4.6, 3.0, 3.0, 3.0, 3.0]),
        ('互补任务条件', 4): [4.0, 4.0, 4.0, 4.6],
        ('预定主要比较', 4): [6.2, 3.2, 4.9, 2.3],
        ('私人资源情境', 3): [8.6, 4.0, 4.0],
        ('评估范围', 5): [4.2, 3.1, 3.1, 3.1, 3.1],
        ('种子', 6): [1.5, 2.6, 2.6, 3.4, 3.4, 3.1],
        ('条件', 6): [3.4, 2.5, 3.0, 2.5, 2.5, 2.7],
        ('配对比较', 3): [6.2, 4.5, 5.9],
        ('来源种子', 5): [2.2, 2.6, 4.6, 2.6, 4.6],
        ('核查对象', 4): [2.6, 1.8, 3.4, 8.8],
        ('记录', 2): [5.0, 11.6],
    }
    if explicit is None:
        explicit = report_widths.get((rows[0][0], count))
    if explicit:
        if len(explicit) != count or any(x <= 0 for x in explicit):
            raise ValueError('table-widths must contain one positive width per column')
        total = sum(explicit)
        if total > CONTENT_WIDTH_CM + .01:
            return [value * CONTENT_WIDTH_CM / total for value in explicit]
        return explicit
    weights, fixed = [], []
    for col in range(count):
        lengths = sorted(text_width(row[col]) for row in rows)
        representative = lengths[min(len(lengths) - 1, math.floor(.8 * len(lengths)))]
        header = text_width(rows[0][col])
        numeric = sum(numeric_cell(row[col]) for row in rows[1:]) >= max(1, .7 * (len(rows) - 1))
        weight = max(5, min(22, .6 * header + .7 * representative))
        weights.append(weight)
        fixed.append(max(1.7, min(3.0, max(header * .14, representative * .14) + .7)) if numeric else None)
    if any(value is None for value in fixed) and any(value is not None for value in fixed):
        flexible_indices = [index for index, value in enumerate(fixed) if value is None]
        remaining = CONTENT_WIDTH_CM - sum(value or 0 for value in fixed)
        base = 1.45
        flexible_weight = sum(weights[index] for index in flexible_indices)
        for index in flexible_indices:
            fixed[index] = base + max(0, remaining - base * len(flexible_indices)) * weights[index] / flexible_weight
        return fixed
    minimum = 1.15 if count >= 7 else 1.45
    remainder = CONTENT_WIDTH_CM - minimum * count
    return [minimum + max(0, remainder) * weight / sum(weights) for weight in weights]


def add_table(doc, lines, widths=None):
    rows = [table_cells(lines[0])] + [table_cells(line) for line in lines[2:]]
    if any(len(row) != len(rows[0]) for row in rows):
        raise ValueError(f'Unequal cell counts in table: {lines[0]}')
    ncols = len(rows[0])
    if ncols > 9:
        raise ValueError('This portrait report supports at most 9 columns; split the table explicitly')
    widths = column_widths(rows, widths)
    size = 9.5 if ncols >= 7 else 10
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    tblpr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    tblpr.append(borders)
    indent = OxmlElement('w:tblInd')
    indent.set(qn('w:w'), '0')
    indent.set(qn('w:type'), 'dxa')
    tblpr.append(indent)
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for r, values in enumerate(rows):
        trpr = table.rows[r]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        trpr.append(no_split)
        if r == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            trpr.append(repeat)
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.width = Cm(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, padding in [('top', 75), ('bottom', 75), ('left', 95), ('right', 95)]:
                elem = OxmlElement(f'w:{side}')
                elem.set(qn('w:w'), str(padding))
                elem.set(qn('w:type'), 'dxa')
                margins.append(elem)
            tcpr.append(margins)
            if r == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E7E6E6')
                shading.set(qn('w:val'), 'clear')
                tcpr.append(shading)
            p = cell.paragraphs[0]
            p_settings(p, size=size, after=1, line=1.16, keep_next=r < len(rows) - 1)
            short_column = all(text_width(row[c]) <= 11 for row in rows)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if short_column else WD_ALIGN_PARAGRAPH.LEFT
            inline(p, value, size=size, bold=r == 0)
    spacer = doc.add_paragraph()
    p_settings(spacer, after=4, line=.4)
    spacer.paragraph_format.line_spacing = Pt(4)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.add_run().font.size = Pt(3)
    return {'rows': len(rows), 'columns': ncols, 'widths_cm': widths}


def resolve_image(target, source):
    target = unquote(target)
    if urlparse(target).scheme not in ('', 'file'):
        raise ValueError('Report figures must be existing local images')
    if target.startswith('file://'):
        target = urlparse(target).path
    path = Path(target)
    return path if path.is_absolute() else (source.parent / path).resolve()


def add_image(doc, line, source, caption=None, width=None):
    match = re.fullmatch(r'!\[(.*)\]\((.*)\)', line.strip())
    if match is None:
        raise ValueError(f'Invalid standalone image: {line}')
    alt, target = match.groups()
    if target.startswith('<') and target.endswith('>'):
        target = target[1:-1]
    image_path = resolve_image(target, source)
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    with Image.open(image_path) as im:
        pixel_width, pixel_height = im.size
    width = min(width or CONTENT_WIDTH_CM, CONTENT_WIDTH_CM)
    height = width * pixel_height / pixel_width
    if height > 16.8:
        width *= 16.8 / height
    p = doc.add_paragraph()
    p_settings(p, before=5, after=3, line=1, keep_next=True)
    p.paragraph_format.line_spacing = 1  # Inline figures require automatic line height.
    p.paragraph_format.keep_together = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(image_path), width=Cm(width))
    shape._inline.docPr.set('descr', alt)
    cp = doc.add_paragraph(style='Caption')
    p_settings(cp, size=9.5, after=8, line=1.18)
    cp.paragraph_format.keep_together = True
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    inline(cp, caption or alt, size=9.5)
    return {'path': str(image_path), 'width_cm': width, 'caption': caption or alt}


def is_special(line):
    return bool(re.match(r'^(?:#{1,6}\s|!\[|>|```|<!--|[-*]\s|\d+[.)]\s)', line)) or line.strip() == '---'


def build(source, output):
    lines = source.read_text(encoding='utf-8').splitlines()
    doc = make_document()
    index = 0
    title_seen = False
    table_width = image_width = None
    manifest = {'source': str(source), 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                'runtime': str(Path(sys.executable)), 'figures': [], 'tables': [], 'headings': [], 'paragraphs': 0}
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if re.fullmatch(r'<!--\s*(?:pagebreak|PAGEBREAK)\s*-->', line):
            doc.add_page_break()
            index += 1
            continue
        if line.startswith('<!--'):
            match = re.fullmatch(r'<!--\s*table-widths:\s*([\d.,\s]+)\s*-->', line)
            if match:
                table_width = [float(x.strip()) for x in match[1].split(',')]
            match = re.fullmatch(r'<!--\s*figure-width:\s*([\d.]+)\s*-->', line)
            if match:
                image_width = float(match[1])
            index += 1
            continue
        if line == '---':
            index += 1
            continue
        heading = re.match(r'^(#{1,6})\s+(.+)$', line)
        if heading:
            level, label = len(heading[1]), heading[2]
            if level == 1 and not title_seen:
                style, size = 'Title', 22
                title_seen = True
            else:
                actual = min(3, max(1, level - 1))
                style, size = f'Heading {actual}', {1: 15, 2: 12.5, 3: 11.5}[actual]
            p = doc.add_paragraph(style=style)
            inline(p, label, size=size, bold=True, heading=True)
            manifest['headings'].append({'level': level, 'text': label, 'style': style})
            index += 1
            continue
        if line.startswith('!['):
            next_index = index + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            caption = lines[next_index].strip() if next_index < len(lines) and CAPTION_RE.match(lines[next_index].strip()) else None
            manifest['figures'].append(add_image(doc, line, source, caption, image_width))
            image_width = None
            index = next_index + 1 if caption else index + 1
            continue
        if index + 1 < len(lines) and '|' in line and re.fullmatch(r'\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*', lines[index + 1]):
            chunk = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and '|' in lines[index] and lines[index].strip():
                chunk.append(lines[index].strip())
                index += 1
            manifest['tables'].append(add_table(doc, chunk, table_width))
            table_width = None
            continue
        if line.startswith('```'):
            chunk = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith('```'):
                chunk.append(lines[index])
                index += 1
            p = doc.add_paragraph()
            p_settings(p, size=9, after=6, line=1.15)
            p.paragraph_format.left_indent = Cm(.3)
            for n, value in enumerate(chunk):
                run = p.add_run(value)
                set_font(run, size=9, latin='Menlo')
                if n < len(chunk) - 1:
                    run.add_break()
            index += 1
            continue
        if line.startswith('>'):
            p = doc.add_paragraph()
            p_settings(p, after=6, line=1.25)
            p.paragraph_format.left_indent, p.paragraph_format.right_indent = Cm(.55), Cm(.3)
            inline(p, line.lstrip('>').strip(), size=10.5)
            index += 1
            continue
        list_match = re.match(r'^([-*]|\d+[.)])\s+(.+)$', line)
        if list_match:
            marker, content = list_match.groups()
            p = doc.add_paragraph()
            p_settings(p, after=3, line=1.27)
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(.55), Cm(-.45)
            inline(p, ('• ' if marker in '-*' else marker + ' ') + content)
            index += 1
            continue
        paragraph = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not is_special(lines[index].strip()) and '|' not in lines[index]:
            paragraph.append(lines[index].strip())
            index += 1
        text = ''.join(paragraph)
        caption = bool(CAPTION_RE.match(text))
        p = doc.add_paragraph(style='Caption' if caption else 'Normal')
        p_settings(p, size=9.5 if caption else 11, after=6 if caption else 5, line=1.18 if caption else 1.32,
                   keep_next=caption and text.startswith(('表', 'Table')))
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if caption else WD_ALIGN_PARAGRAPH.JUSTIFY
        inline(p, text, size=9.5 if caption else 11)
        manifest['paragraphs'] += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    manifest['output'] = str(output)
    manifest['output_sha256'] = hashlib.sha256(output.read_bytes()).hexdigest()
    output.with_suffix('.build.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=HERE / '研究报告.md')
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if Path(sys.executable).resolve() != BUNDLED_PYTHON.resolve():
        raise RuntimeError(f'Use the bundled runtime: {BUNDLED_PYTHON}')
    result = build(args.source.resolve(), args.output.resolve())
    print(json.dumps({key: result[key] for key in ['source', 'output', 'paragraphs']}, ensure_ascii=False))
    print(f"Figures: {len(result['figures'])}; tables: {len(result['tables'])}; headings: {len(result['headings'])}")


if __name__ == '__main__':
    main()
